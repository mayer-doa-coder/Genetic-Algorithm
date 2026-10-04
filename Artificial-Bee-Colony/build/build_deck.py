"""Build Ackley_ABC_Presentation.pptx.

The deck is built *on top of* the bee template, so it inherits that file's
background, its embedded fonts (More Sugar, DM Sans) and - most importantly -
its illustrations. Every honeycomb, bee, hive and honey drip in the finished
deck is the original template artwork, copied shape-for-shape by `deco()`.
"""

from pathlib import Path
from copy import deepcopy
import json

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

from content import (SLIDES, QA, DATE, COURSE, TEACHER, TEACHER_ROLE, IDS,
                     CONTROL_PARAMETERS, GA_TO_ABC, WHEEL_TABLE, ABLATION,
                     COMPARISON, BEE_ROLES, TITLE_WIDTH_AT_80)

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FIGURES = ROOT / "figures"
ASSETS = ROOT / "review" / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)
TEMPLATE = ROOT.parent / "Yellow and Brown Illustrated Bee Group Project Presentation.pptx"

# --------------------------------------------------------------- palette
BG      = "F5F5F5"   # the template background
BROWN   = "6F251B"   # the template's only text colour
AMBER   = "FFC62A"   # honeycomb yellow
DEEP    = "E37E03"   # burnt honey
HIVE    = "CB6C29"   # hive brown
DARK    = "803B1B"   # darkest wood
WHITE   = "FFFFFF"
CREAM   = "FDF6E6"   # soft card fill
SAND    = "F3E7D2"   # alternate table row
LINE    = "E3D6CC"

DISPLAY = "More Sugar"     # embedded in the template
BODY    = "DM Sans"
BODYB   = "DM Sans Bold"

I = Inches


# =====================================================================
# Open the template, harvest its artwork, then empty it
# =====================================================================
prs = Presentation(TEMPLATE)

# (template slide index, shape name) -> our component name
COMPONENTS = {
    "honeycomb":      (1, "Freeform 2"),    # big honeycomb cluster
    "hive_bees":      (1, "Freeform 4"),    # beehive with four bees
    "bee":            (1, "Freeform 5"),    # single yellow line-art bee
    "hex_outline":    (2, "Freeform 2"),    # three hexagon outlines
    "hanging_hive":   (3, "Freeform 3"),    # hive hanging from a branch
    "honey_pot":      (4, "Freeform 4"),    # pot with dipper and bees
    "honey_drip":     (4, "Freeform 5"),    # dripping honey
    "hive_branch":    (5, "Freeform 3"),    # hive on a branch, three bees
    "honey_jar":      (6, "Freeform 2"),    # jar with a bee on the label
    "comb_glow":      (7, "Freeform 2"),    # glowing honeycomb
    "flower_vine":    (8, "Freeform 6"),    # vine, leaves and a bee
    "hive_flowers":   (10, "Freeform 10"),  # hive surrounded by flowers
}

SOURCES = {}
for name, (slide_no, shape_name) in COMPONENTS.items():
    slide = prs.slides[slide_no - 1]
    shape = next(sh for sh in slide.shapes if sh.name == shape_name)
    rels = {}
    for node in shape._element.iter():
        rid = node.get(qn("r:embed"))
        if rid and rid not in rels:
            rel = slide.part.rels[rid]
            rels[rid] = (rel.target_part, rel.reltype)
    SOURCES[name] = (deepcopy(shape._element), rels)

for sld_id in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)

BLANK = prs.slide_layouts[6]


def prime_components(slide):
    """Re-attach every harvested image part to the package straight away.

    Clearing the template's slides leaves its illustrations unreachable, so
    python-pptx would hand out partnames such as /ppt/media/image8.png to our
    own figures - and then `deco()` would reintroduce the original image8 under
    the same name, producing a file PowerPoint refuses to open. Relating all of
    them to the first slide before any picture is added keeps the names unique.
    """
    for _, rels in SOURCES.values():
        for part, reltype in rels.values():
            slide.part.relate_to(part, reltype)


# =====================================================================
# Drawing helpers
# =====================================================================
def color(value):
    return RGBColor.from_string(value)


def deco(s, name, x, y, w, h, flip_h=False, rotation=None):
    """Place one of the template's original illustrations."""
    element, rels = SOURCES[name]
    element = deepcopy(element)
    for node in element.iter():
        rid = node.get(qn("r:embed"))
        if rid in rels:
            part, reltype = rels[rid]
            node.set(qn("r:embed"), s.part.relate_to(part, reltype))
    s.shapes._spTree.insert_element_before(element, "p:extLst")
    shape = s.shapes[-1]
    shape._element.xpath(".//p:cNvPr")[0].set("id", str(s.shapes._next_shape_id))
    shape.left, shape.top, shape.width, shape.height = I(x), I(y), I(w), I(h)
    shape.name = f"Template artwork - {name}"
    if flip_h:
        shape._element.spPr.xfrm.set("flipH", "1")
    if rotation is not None:
        shape.rotation = rotation
    return shape


def card(s, x, y, w, h, fill=WHITE, line=None, radius=True, name=None):
    shape = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        I(x), I(y), I(w), I(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1.4)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = min(0.09, 0.5 * min(w, h) / max(w, h) * 0.6 + 0.04)
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def text(s, body, x, y, w, h, size=26, font=BODY, c=BROWN, align=None,
         bold=False, anchor=None, spacing=1.0, name=None):
    shape = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    if anchor is not None:
        frame.vertical_anchor = anchor
    for i, line in enumerate(str(body).split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(c)
        p.space_after = Pt(size * 0.22)
        p.space_before = Pt(0)
        p.line_spacing = spacing
        if align is not None:
            p.alignment = align
    if name:
        shape.name = name
    return shape


def heading(s, body, x=1.12, y=1.30, w=14.0, h=1.45, size=80, align=PP_ALIGN.LEFT,
            name="Slide title"):
    return text(s, body, x, y, w, h, size, DISPLAY, BROWN, align, name=name)


def subhead(s, body, x, y, w, h=0.55, size=30, c=BROWN, align=PP_ALIGN.LEFT, name=None):
    return text(s, body, x, y, w, h, size, BODYB, c, align, bold=True, name=name)


def label(s, body, x, y, w=9.0, size=19, c=HIVE):
    return text(s, body, x, y, w, 0.38, size, BODYB, c, bold=True)


def hexbadge(s, number, x, y, size=1.22, fill=BROWN, fg=WHITE, name=None):
    """The template's numbered hexagon: a flat hexagon with a short label."""
    shape = s.shapes.add_shape(MSO_SHAPE.HEXAGON, I(x), I(y), I(size), I(size * 0.86))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.text = str(number)
    p.alignment = PP_ALIGN.CENTER
    p.font.name = DISPLAY
    p.font.size = Pt(34 if len(str(number)) <= 2 else 26)
    p.font.color.rgb = color(fg)
    return shape


def bullet(s, number, title, bodytext, x, y, w, badge=BROWN, title_size=31,
           body_size=24, name=None):
    """A hexagon badge with a bold line and an explanation next to it."""
    hexbadge(s, number, x, y + 0.03, 1.22, badge, name=(name + "_badge") if name else None)
    subhead(s, title, x + 1.62, y - 0.02, w - 1.7, 0.55, title_size,
            name=(name + "_title") if name else None)
    if bodytext:
        text(s, bodytext, x + 1.62, y + 0.60, w - 1.7, 0.9, body_size, BODY, BROWN,
             PP_ALIGN.LEFT, name=(name + "_body") if name else None)


def arrow(s, x, y, w=0.72, h=0.42, c=DEEP, rotation=0):
    shape = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, I(x), I(y), I(w), I(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(c)
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.rotation = rotation
    return shape


def rule(s, x1, y1, x2, y2, c=LINE, width=1.4):
    shape = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    shape.line.color.rgb = color(c)
    shape.line.width = Pt(width)
    return shape


def pic(s, name, x, y, w, h):
    """Place a generated figure, fitted inside the given box."""
    path = FIGURES / name
    iw, ih = Image.open(path).size
    ratio = min(w / iw, h / ih)
    pw, ph = iw * ratio, ih * ratio
    shape = s.shapes.add_picture(str(path), I(x + (w - pw) / 2), I(y + (h - ph) / 2),
                                 width=I(pw), height=I(ph))
    shape.name = name
    return shape


def equation(s, formula, x, y, w, h, key, size=27):
    """Typeset a formula with matplotlib, in the template's brown."""
    fig = plt.figure(figsize=(16, 2))
    fig.patch.set_alpha(0)
    fig.text(0.5, 0.5, "$" + formula + "$", ha="center", va="center",
             fontsize=size, color="#" + BROWN)
    path = ASSETS / (key + ".png")
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=0.08, transparent=True)
    plt.close(fig)
    iw, ih = Image.open(path).size
    ratio = min(w / iw, h / ih)
    shape = s.shapes.add_picture(str(path), I(x + (w - iw * ratio) / 2),
                                 I(y + (h - ih * ratio) / 2),
                                 width=I(iw * ratio), height=I(ih * ratio))
    shape.name = key
    shape._element.xpath(".//p:cNvPr")[0].set("descr", formula)
    return shape


def table(s, headers, rows, x, y, widths, row_height=0.72, size=22,
          highlight=None, warn=None):
    frame = s.shapes.add_table(len(rows) + 1, len(headers), I(x), I(y),
                               I(sum(widths)), I(row_height * (len(rows) + 1)))
    t = frame.table
    for j, w in enumerate(widths):
        t.columns[j].width = I(w)
    for i, row in enumerate([headers] + rows):
        t.rows[i].height = I(row_height)
        for j, value in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(value)
            cell.margin_left = I(0.20)
            cell.margin_right = I(0.14)
            cell.margin_top = I(0.06)
            cell.margin_bottom = I(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                fill, fg, bold = BROWN, WHITE, True
            elif i == highlight:
                fill, fg, bold = AMBER, BROWN, True
            elif i == warn:
                fill, fg, bold = HIVE, WHITE, True
            else:
                fill, fg, bold = (WHITE if i % 2 else SAND), BROWN, False
            cell.fill.fore_color.rgb = color(fill)
            for p in cell.text_frame.paragraphs:
                p.font.name = BODYB if bold else BODY
                p.font.size = Pt(size)
                p.font.bold = bold
                p.font.color.rgb = color(fg)
                if j > 0 and len(headers) > 2:
                    p.alignment = PP_ALIGN.LEFT
    return frame


def takeaway(s, message, y=9.72, fill=AMBER, c=BROWN):
    card(s, 1.12, y, 17.76, 0.82, fill)
    text(s, message, 1.52, y + 0.17, 17.0, 0.5, 26, BODYB, c, PP_ALIGN.LEFT, bold=True)


def stat(s, value, caption, x, y, w, h=1.72, fill=WHITE, value_size=42, c=BROWN):
    card(s, x, y, w, h, fill)
    text(s, value, x + 0.30, y + 0.18, w - 0.6, 0.8, value_size, DISPLAY, c)
    text(s, caption, x + 0.30, y + h - 0.62, w - 0.6, 0.45, 21, BODY, BROWN)


# =====================================================================
# Slide shell
# =====================================================================
TITLE_MAX_WIDTH = 15.40   # leaves the top-right corner free for an illustration


def title_size_for(n):
    """Pick the largest size that still keeps the title on one line."""
    return min(80, int(80 * TITLE_MAX_WIDTH / TITLE_WIDTH_AT_80[n - 1]))


def new(n, title_y=1.22, show_title=True):
    data = SLIDES[n - 1]
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(BG)
    for shape in list(s.shapes):
        s.shapes._spTree.remove(shape._element)

    if show_title:
        label(s, data["section"], 1.16, title_y - 0.46, 12.0)
        heading(s, data["title"], 1.12, title_y, TITLE_MAX_WIDTH + 0.6, 1.40,
                title_size_for(n))

    # footer: a small amber hexagon with the slide number
    hexbadge(s, f"{n:02d}", 18.08, 10.40, 0.90, AMBER, BROWN, name="Slide number")
    text(s, DATE, 1.12, 10.58, 6.0, 0.35, 16, BODY, HIVE, name="Presentation date")
    text(s, "ACKLEY  x  ARTIFICIAL BEE COLONY", 6.5, 10.58, 7.0, 0.35, 15,
         BODYB, HIVE, PP_ALIGN.CENTER)

    s.notes_slide.notes_text_frame.text = (
        f"SLIDE {n} - {data['title']}\n"
        f"Target: {data['seconds']} seconds\n\n"
        f"{data['say']}\n\n"
        f"VISUAL / DELIVERY\n{data['figure']}\n\n"
        "Source: Ackley_ABC_Solution.ipynb; every number reproduced by "
        "review/experiments.py."
    )
    return s


# =====================================================================
# 01  Title
# =====================================================================
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = color(BG)
for shape in list(s.shapes):
    s.shapes._spTree.remove(shape._element)
prime_components(s)

deco(s, "honeycomb", -0.91, -1.53, 7.65, 5.32)
deco(s, "honeycomb", 12.46, 7.71, 8.72, 6.07)
deco(s, "hive_bees", 15.49, 4.50, 7.34, 5.96)
for bx, by in [(17.72, 1.12), (0.79, 9.32), (12.56, -0.52), (-0.13, 6.33), (19.10, -0.52)]:
    deco(s, "bee", bx, by, 1.16, 1.03)

label(s, COURSE.upper() + "   -   GROUP PRESENTATION", 1.12, 1.95, 14.0, 21, DEEP)
text(s, "Artificial\nBee Colony", 1.12, 2.50, 14.5, 4.10, 98, DISPLAY, BROWN)
text(s, "Finding the global minimum of the 2-D Ackley function",
     1.12, 6.30, 13.6, 0.70, 33, BODY, HIVE)
rule(s, 1.12, 7.25, 14.2, 7.25, LINE, 2.0)

label(s, "SUBMITTED TO", 1.12, 7.56, 8.0)
text(s, TEACHER, 1.12, 7.98, 9.0, 0.62, 36, BODYB, BROWN, bold=True)
text(s, TEACHER_ROLE, 1.12, 8.62, 10.0, 0.45, 25, BODY, HIVE)
label(s, "SUBMITTED BY   -   STUDENT IDs", 1.12, 9.26, 9.0)
text(s, "   ".join(IDS[:3]) + "\n" + "   ".join(IDS[3:]),
     1.12, 9.66, 11.0, 1.05, 26, BODY, BROWN)

s.notes_slide.notes_text_frame.text = (
    f"SLIDE 1 - {SLIDES[0]['title']}\nTarget: {SLIDES[0]['seconds']} seconds\n\n"
    f"{SLIDES[0]['say']}\n\nVISUAL / DELIVERY\n{SLIDES[0]['figure']}")

# =====================================================================
# 02  Find the lowest point
# =====================================================================
s = new(2)
deco(s, "honeycomb", -1.70, 7.55, 6.60, 4.60)
deco(s, "bee", 17.30, 0.42, 1.10, 0.98)

bullet(s, "01", "Choose two numbers", "Pick x₁ and x₂ so that f is as small as possible.",
       1.12, 3.10, 9.3)
bullet(s, "02", "Stay inside the square", "Both coordinates must satisfy  −5 ≤ x ≤ +5.",
       1.12, 4.90, 9.3, DEEP)
bullet(s, "03", "The answer we are aiming at", "x* = (0, 0)  and  f(x*) = 0.",
       1.12, 6.70, 9.3, HIVE)

# a plain coordinate square, drawn natively
card(s, 11.30, 2.88, 7.58, 6.55, WHITE)
left, top, side = 12.20, 3.50, 5.00
for k in range(6):
    offset = side * k / 5
    rule(s, left + offset, top, left + offset, top + side, "EFE6E0", 0.9)
    rule(s, left, top + offset, left + side, top + offset, "EFE6E0", 0.9)
rule(s, left - 0.45, top + side / 2, left + side + 0.60, top + side / 2, HIVE, 2.2)
rule(s, left + side / 2, top - 0.45, left + side / 2, top + side + 0.55, HIVE, 2.2)
dot = s.shapes.add_shape(MSO_SHAPE.OVAL, I(left + side / 2 - 0.15),
                         I(top + side / 2 - 0.15), I(0.30), I(0.30))
dot.fill.solid(); dot.fill.fore_color.rgb = color(BROWN); dot.line.fill.background()
dot.shadow.inherit = False
text(s, "(0, 0)", left + side / 2 + 0.26, top + side / 2 + 0.10, 2.2, 0.5, 27, BODYB, BROWN,
     bold=True)
text(s, "−5", left - 0.30, top + side + 0.12, 0.9, 0.4, 21, BODY, HIVE)
text(s, "+5", left + side - 0.45, top + side + 0.12, 0.9, 0.4, 21, BODY, HIVE)
text(s, "x₁", left + side + 0.72, top + side / 2 - 0.22, 0.8, 0.5, 26, BODYB, HIVE, bold=True)
text(s, "x₂", left + side / 2 - 0.70, top - 0.72, 0.9, 0.5, 26, BODYB, HIVE, bold=True)
text(s, "every point in the square is allowed", 11.55, 8.92, 7.1, 0.4, 21, BODY, HIVE,
     PP_ALIGN.CENTER)

takeaway(s, "The bees only taste the points they choose. Nobody tells them where (0, 0) is.")

# =====================================================================
# 03  The Ackley function in 2-D
# =====================================================================
s = new(3)
deco(s, "honey_drip", 16.9, -1.05, 3.6, 2.78)

label(s, "START WITH n VARIABLES", 1.12, 3.05, 12.0)
card(s, 1.12, 3.52, 17.76, 1.80, WHITE)
equation(s, r"f(\mathbf{x})=-a\exp\!\left(-b\sqrt{\frac{1}{n}\sum_{i=1}^{n}x_i^2}\right)"
            r"-\exp\!\left(\frac{1}{n}\sum_{i=1}^{n}\cos(cx_i)\right)+a+e",
         1.45, 3.72, 17.1, 1.40, "general_ackley")

for caption, x in [("Set n = 2", 1.50), ("Expand the two sums", 7.30),
                   ("Use a = 20, b = 0.2, c = 2π", 13.10)]:
    text(s, caption, x, 5.62, 4.80, 0.6, 26, BODYB, DEEP, PP_ALIGN.CENTER, bold=True)
arrow(s, 6.45, 5.72, 0.70, 0.40)
arrow(s, 12.25, 5.72, 0.70, 0.40)

label(s, "THE FUNCTION IN OUR CODE", 1.12, 6.55, 12.0)
card(s, 1.12, 7.02, 17.76, 2.00, CREAM)
equation(s, r"f(x_1,x_2)=-20\exp\!\left(-0.2\sqrt{\frac{x_1^2+x_2^2}{2}}\right)"
            r"-\exp\!\left(\frac{\cos(2\pi x_1)+\cos(2\pi x_2)}{2}\right)+20+e",
         1.45, 7.24, 17.1, 1.56, "two_dimensional_ackley")

takeaway(s, "Check it at the origin:    −20 − e + 20 + e  =  0")

# =====================================================================
# 04  One deep minimum, many traps
# =====================================================================
s = new(4)
deco(s, "bee", 17.40, 0.40, 1.10, 0.98)

card(s, 1.12, 2.88, 12.10, 4.95, WHITE)
pic(s, "01_ackley_landscape.png", 1.24, 2.98, 11.86, 4.75)
card(s, 1.12, 7.98, 12.10, 2.32, WHITE)
pic(s, "02_ackley_slice.png", 1.24, 8.08, 11.86, 2.12)

bullet(s, "01", "One wide funnel", "The smooth part slopes towards the centre.",
       13.58, 2.95, 5.30, BROWN, 27, 21)
bullet(s, "02", "Many small traps", "The cosine part adds dips everywhere.",
       13.58, 4.75, 5.30, DEEP, 27, 21)
bullet(s, "03", "Too big to scan", "A 10⁻⁶ grid would need ~10¹⁴ points.",
       13.58, 6.55, 5.30, HIVE, 27, 21)

card(s, 13.58, 8.40, 5.30, 1.90, CREAM)
text(s, "A plain downhill search\ncan stop in any of the dips.",
     13.90, 8.78, 4.70, 1.3, 25, BODYB, BROWN, bold=True)

# =====================================================================
# 05  A colony instead of a population
# =====================================================================
s = new(5)
deco(s, "honeycomb", -2.20, 8.55, 6.60, 4.60)

label(s, "BEE LANGUAGE → OUR PROGRAM", 1.12, 2.90, 12.0)
table(s, ["In the hive", "In the program"],
      [["Food source", "one answer [x₁, x₂]"],
       ["Nectar amount", "fitness, bigger is better"],
       ["Employed bee", "improves its own source"],
       ["Onlooker bee", "helps the richest sources"],
       ["Scout bee", "restarts a stuck source"]],
      1.12, 3.36, [3.70, 5.05], 0.80, 24)
card(s, 1.12, 8.32, 8.68, 1.12, CREAM)
text(s, "Published by Dervis Karaboga in 2005,\nafter the foraging behaviour of honey bees.",
     1.45, 8.52, 8.1, 1.0, 23, BODY, BROWN)

card(s, 10.20, 2.88, 8.68, 6.55, WHITE)
pic(s, "03_initial_sources.png", 10.38, 2.98, 8.32, 6.35)

takeaway(s, "25 food sources explore at once. Better sources attract more of the colony's effort.")

# =====================================================================
# 06  Three bees, three jobs
# =====================================================================
s = new(6)
deco(s, "hive_branch", 14.1, -2.30, 6.3, 4.3)

PANEL_COLOURS = [AMBER, DEEP, HIVE]
for k, (title, sub, bodytext, purpose) in enumerate(BEE_ROLES):
    x = 1.12 + k * 6.05
    accent = PANEL_COLOURS[k]
    reveal = f"Reveal_{k + 1}"
    card(s, x, 2.92, 5.65, 6.50, WHITE, name=reveal)
    card(s, x, 2.92, 5.65, 1.02, accent, name=reveal + "_head")
    text(s, title, x + 0.34, 3.16, 5.0, 0.6, 26, BODYB,
         BROWN if accent == AMBER else WHITE, bold=True, name=reveal + "_title")
    text(s, sub, x + 0.34, 4.14, 5.0, 0.5, 20, BODY, HIVE, name=reveal + "_sub")
    text(s, bodytext, x + 0.34, 4.78, 5.05, 1.7, 24, BODY, BROWN, name=reveal + "_body")
    deco(s, "bee", x + 2.05, 6.60, 1.55, 1.38)
    text(s, purpose, x + 0.34, 8.42, 5.0, 0.7, 30, DISPLAY,
         accent if accent != AMBER else DEEP, PP_ALIGN.CENTER, name=reveal + "_purpose")
    if k < 2:
        arrow(s, x + 5.76, 5.80, 0.45, 0.42)

takeaway(s, "Exploitation, guided search and exploration — three jobs, repeated every cycle.")

# =====================================================================
# 07  The settings we used
# =====================================================================
s = new(7)
deco(s, "honey_jar", 17.10, -0.45, 2.3, 2.16)

label(s, "THE SIX CONTROL PARAMETERS (SLIDE 20)", 1.12, 2.90, 9.0)
table(s, ["Parameter", "Value"], CONTROL_PARAMETERS, 1.12, 3.36, [5.30, 3.10], 0.78, 24)
card(s, 1.12, 8.14, 8.40, 1.62, CREAM)
text(s, "limit = SN × D = 25 × 2 = 50\nSeed 7, so the run repeats exactly.",
     1.45, 8.46, 7.8, 1.1, 24, BODYB, BROWN, bold=True)

label(s, "THE SAME BUDGET, DIFFERENT MACHINERY", 10.48, 2.90, 9.0, 19, DEEP)
table(s, ["Genetic Algorithm", "Artificial Bee Colony"], GA_TO_ABC,
      10.48, 3.36, [4.20, 4.20], 0.68, 23)
card(s, 10.48, 8.14, 8.40, 1.62, AMBER)
text(s, "Two blank rows: ABC has no crossover\nrate and no mutation rate to choose.",
     10.80, 8.46, 7.8, 1.1, 24, BODYB, BROWN, bold=True)

# =====================================================================
# 08  The ABC cycle
# =====================================================================
s = new(8)
deco(s, "bee", 17.40, 0.40, 1.10, 0.98)
deco(s, "honeycomb", 14.90, -3.95, 7.20, 5.00)


def flowbox(x, y, w, h, title, bodytext, fill=WHITE, accent=None, name=None,
            title_size=27, body_size=21):
    card(s, x, y, w, h, fill, name=name)
    if accent:
        card(s, x, y, 0.16, h, accent, radius=False)
    text(s, title, x + 0.42, y + 0.18, w - 0.75, 0.52, title_size, BODYB, BROWN,
         bold=True, name=(name + "_title") if name else None)
    if bodytext:
        text(s, bodytext, x + 0.42, y + 0.78, w - 0.75, 0.6, body_size, BODY, HIVE,
             name=(name + "_body") if name else None)


flowbox(1.12, 3.00, 7.10, 1.52, "1   Assign the parameters",
        "colony 50, SN 25, limit 50, 100 cycles", CREAM, BROWN)
flowbox(1.12, 4.82, 7.10, 1.52, "2   Initialise 25 sources",
        "random points, evaluate, remember the best", CREAM, BROWN)
rule(s, 4.67, 4.52, 4.67, 4.82, DEEP, 2.4)
arrow(s, 8.32, 5.22, 0.72, 0.42)

card(s, 9.28, 2.86, 9.60, 6.42, WHITE)
text(s, "Step 3  —  repeat every cycle", 9.62, 3.02, 8.9, 0.55, 28, DISPLAY, DEEP)
flowbox(9.62, 3.72, 8.92, 1.26, "Employed bee phase",
        "25 bees search next to their own source", SAND, AMBER, "Reveal_1", 25, 20)
flowbox(9.62, 5.12, 8.92, 1.26, "Onlooker bee phase",
        "25 roulette-wheel visits to the richer sources", SAND, DEEP, "Reveal_2", 25, 20)
flowbox(9.62, 6.52, 8.92, 1.26, "Memorise the best solution",
        "done BEFORE the scouts run", AMBER, BROWN, "Reveal_3", 25, 20)
flowbox(9.62, 7.92, 8.92, 1.26, "Scout bee phase",
        "replace any source stuck for 50 tries", SAND, HIVE, "Reveal_4", 25, 20)
for y in (4.98, 6.38, 7.78):
    rule(s, 14.08, y, 14.08, y + 0.14, DEEP, 2.4)

flowbox(1.12, 6.64, 7.10, 1.52, "4   Stop and report",
        "after 100 cycles, or when progress stalls", CREAM, BROWN)
text(s, "↺   back to step 3", 1.50, 8.46, 6.5, 0.55, 27, DISPLAY, DEEP)

takeaway(s, "Memorising the best before the scouts run is this algorithm's elitism.")

# =====================================================================
# 09  Every solution is two real numbers
# =====================================================================
s = new(9)
deco(s, "comb_glow", 16.90, -1.05, 4.60, 3.00)

label(s, "ONE FOOD SOURCE", 1.12, 2.90, 12.0)
for x, symbol, caption, accent in [(2.65, "x₁", "COORDINATE 1", DEEP),
                                   (10.20, "x₂", "COORDINATE 2", HIVE)]:
    label(s, caption, x, 3.34, 6.5, 19, accent)
    card(s, x, 3.76, 6.55, 2.25, WHITE)
    text(s, symbol, x, 4.02, 6.55, 1.7, 78, DISPLAY, accent, PP_ALIGN.CENTER)
text(s, "both inside  [ −5 ,  +5 ]", 2.65, 6.20, 14.1, 0.6, 28, BODYB, BROWN,
     PP_ALIGN.CENTER, bold=True)

label(s, "INITIALISE  —  AND THE SCOUTS REUSE THE SAME LINE", 1.12, 7.02, 14.0)
card(s, 1.12, 7.44, 11.20, 2.00, CREAM)
equation(s, r"x_{ij} = x_{\min,j} + \mathrm{rand}(0,1)\,\left(x_{\max,j} - x_{\min,j}\right)",
         1.45, 7.66, 10.6, 1.56, "initialisation")
card(s, 12.72, 7.44, 6.16, 2.00, WHITE)
text(s, "Example sources", 13.05, 7.72, 5.5, 0.4, 20, BODYB, HIVE, bold=True)
text(s, "[ −4.0 ,  2.5 ]\n[ −1.5 ,  −1.0 ]",
     13.05, 8.20, 5.6, 1.1, 26, BODYB, BROWN, bold=True)

takeaway(s, "Value encoding: the food source stores the two coordinates directly.")

# =====================================================================
# 10  How a bee searches nearby
# =====================================================================
s = new(10)
deco(s, "bee", 17.40, 0.40, 1.10, 0.98)

card(s, 1.12, 2.88, 17.76, 1.66, AMBER)
equation(s, r"v_{ij} \;=\; x_{ij} \;+\; \varphi_{ij}\,\left(x_{ij} - x_{kj}\right)",
         1.45, 2.99, 17.1, 1.44, "neighbour", size=30)

for caption, x in [("φ = rand(−1, 1)", 1.12), ("k ≠ i, another source", 7.04),
                   ("j = one random coordinate", 12.96)]:
    text(s, caption, x, 4.72, 5.80, 0.5, 24, BODYB, HIVE, PP_ALIGN.CENTER, bold=True)

card(s, 1.12, 5.34, 17.76, 4.10, WHITE)
pic(s, "07_neighbour_step.png", 1.26, 5.42, 17.48, 3.94)

takeaway(s, "The step is a fraction of the gap to another bee — so it shrinks by itself.")

# =====================================================================
# 11  Keep the better one
# =====================================================================
s = new(11)
deco(s, "honey_pot", 17.35, 5.60, 2.95, 4.55)

card(s, 1.12, 2.92, 7.10, 1.40, WHITE)
text(s, "old source   xᵢ", 1.45, 3.12, 6.4, 0.9, 30, BODYB, BROWN, bold=True)
arrow(s, 8.36, 3.22, 0.80, 0.46)
card(s, 9.30, 2.92, 7.10, 1.40, CREAM)
text(s, "trial point   vᵢ", 9.63, 3.12, 6.4, 0.9, 30, BODYB, DEEP, bold=True)

text(s, "compare  f(xᵢ)  with  f(vᵢ)", 1.12, 4.62, 15.3, 0.7, 30, DISPLAY, BROWN,
     PP_ALIGN.CENTER)

card(s, 1.12, 5.58, 7.52, 3.85, AMBER)
text(s, "vᵢ is better", 1.52, 5.84, 6.7, 0.6, 31, BODYB, BROWN, bold=True)
text(s, "Replace the source with vᵢ.\nReset its trial counter to 0.",
     1.52, 6.62, 6.7, 1.5, 25, BODY, BROWN)
text(s, "trial → 0", 1.52, 8.46, 6.7, 0.7, 32, DISPLAY, BROWN)

card(s, 9.30, 5.58, 7.52, 3.85, WHITE)
text(s, "vᵢ is worse", 9.70, 5.84, 6.7, 0.6, 31, BODYB, HIVE, bold=True)
text(s, "Keep the old source.\nAdd one to its trial counter.",
     9.70, 6.62, 6.7, 1.5, 25, BODY, BROWN)
text(s, "trial → trial + 1", 9.70, 8.46, 6.7, 0.7, 32, DISPLAY, HIVE)

takeaway(s, "Greedy selection: a food source can never get worse. The counter is what scouts watch.")

# =====================================================================
# 12  Better sources get larger slices
# =====================================================================
s = new(12)

label(s, "LOWER f  →  MORE NECTAR", 1.12, 2.90, 10.0)
card(s, 1.12, 3.32, 10.05, 1.95, CREAM)
equation(s, r"fit_i=\frac{1}{1+f_i}\quad (f_i\geq 0),\qquad "
            r"P_i=\frac{fit_i}{\sum_{m=1}^{25} fit_m}",
         1.40, 3.50, 9.5, 1.60, "nectar")
text(s, "Ackley is never negative, so the second branch of\nthe standard rule is never used here.",
     1.12, 5.42, 10.0, 0.95, 20, BODY, HIVE)

table(s, ["Source", "f", "Share of wheel"], WHEEL_TABLE,
      1.12, 6.26, [3.55, 2.60, 3.90], 0.66, 23, highlight=1)

card(s, 11.55, 2.86, 7.33, 6.58, WHITE)
pic(s, "04_roulette_pie_cycle0.png", 11.68, 2.94, 7.07, 6.42)

takeaway(s, "Each of the 25 onlookers spins this wheel once. Even the poorest source keeps 3 %.")

# =====================================================================
# 13  Something the slides do not mention
# =====================================================================
s = new(13)
deco(s, "bee", 17.40, 0.40, 1.10, 0.98)

card(s, 1.12, 2.88, 9.30, 4.70, WHITE)
pic(s, "05_roulette_early_vs_late.png", 1.24, 2.96, 9.06, 4.54)
card(s, 10.78, 2.88, 8.10, 4.70, WHITE)
pic(s, "06_wheel_pressure.png", 10.88, 2.96, 7.90, 4.54)

bullet(s, "01", "The nectar rule has a ceiling",
       "fit = 1/(1+f) can never be larger than 1.", 1.12, 7.90, 8.9, DEEP, 27, 22)
bullet(s, "02", "So the wheel goes flat",
       "From cycle 62 every slice is exactly 4.00 %.", 10.00, 7.90, 8.9, HIVE, 27, 22)

takeaway(s, "We measured it instead of guessing. The next slide shows what it actually costs.")

# =====================================================================
# 14  Which part of ABC actually matters?
# =====================================================================
s = new(14)

text(s, "Every row changes one thing and runs the same 40 seeds (0–39).",
     1.12, 2.80, 17.0, 0.5, 23, BODY, HIVE)
table(s, ["What we changed", "median f", "worst f", "reached 10⁻⁶", "scouts sent"],
      ABLATION, 1.12, 3.30, [5.40, 3.10, 3.10, 3.06, 3.10], 0.62, 22,
      highlight=1, warn=6)

for x, title, bodytext, fill, fg in [
        (1.12, "Onlookers do the precision work",
         "Remove them and the median is 500,000x worse.", AMBER, BROWN),
        (7.18, "Scouts never fired here",
         "Turning them off changed nothing at all.", WHITE, BROWN),
        (13.24, "A small limit is destructive",
         "limit = 5 threw away good sources: 0 / 40.", HIVE, WHITE)]:
    card(s, x, 8.52, 5.64, 1.76, fill)
    text(s, title, x + 0.32, 8.72, 5.05, 0.55, 23, BODYB, fg, bold=True)
    text(s, bodytext, x + 0.32, 9.38, 5.05, 0.80, 20, BODY, fg)

# =====================================================================
# 15  Our result is very close to (0, 0)
# =====================================================================
s = new(15)

# the path figure is wide and short, so it gets the left two thirds
card(s, 1.12, 2.88, 11.55, 6.56, WHITE)
pic(s, "09_best_solution_path.png", 1.24, 2.98, 11.31, 6.36)

stat(s, "6.541 × 10⁻¹³", "best Ackley value, seed 7", 13.00, 2.88, 5.88, 2.00, AMBER, 40)
stat(s, "99 / 100", "cycle the best was found", 13.00, 5.16, 5.88, 2.00, WHITE, 40)
stat(s, "5,025", "Ackley evaluations used, 0 scouts", 13.00, 7.44, 5.88, 2.00, WHITE, 40)

card(s, 1.12, 9.72, 11.55, 0.82, CREAM)
text(s, "x₁ = +2.278 × 10⁻¹³      x₂ = −3.938 × 10⁻¹⁴      distance = 2.312 × 10⁻¹³",
     1.35, 9.89, 11.1, 0.5, 23, BODYB, BROWN, PP_ALIGN.CENTER, bold=True)
card(s, 13.00, 9.72, 5.88, 0.82, AMBER)
text(s, "random search: 0.769", 13.20, 9.89, 5.5, 0.5, 23, BODYB, BROWN,
     PP_ALIGN.CENTER, bold=True)

# =====================================================================
# 16  The colony closing in
# =====================================================================
s = new(16)

# the snapshot grid is nearly square, so it takes the left column
card(s, 1.12, 2.88, 9.05, 6.56, WHITE)
pic(s, "10_source_snapshots.png", 1.22, 2.96, 8.85, 6.40)
text(s, "the colony contracts", 1.12, 9.46, 9.05, 0.4, 20, BODYB, HIVE, PP_ALIGN.CENTER,
     bold=True)

card(s, 10.45, 2.88, 8.43, 3.30, WHITE)
pic(s, "08_convergence_and_work.png", 10.55, 2.96, 8.23, 3.14)

card(s, 10.45, 6.42, 8.43, 3.02, WHITE)
for k, (title, bodytext, accent) in enumerate([
        ("The best value never rises",
         "it is memorised at the end of every cycle", BROWN),
        ("Both phases keep working",
         "improvements continue right up to cycle 100", DEEP),
        ("The trial counter peaked at 23",
         "the limit of 50 was never reached, so no scout flew", HIVE)]):
    y = 6.62 + k * 0.95
    card(s, 10.78, y + 0.06, 0.14, 0.62, accent, radius=False)
    text(s, title, 11.12, y, 7.5, 0.44, 24, BODYB, BROWN, bold=True)
    text(s, bodytext, 11.12, y + 0.44, 7.5, 0.40, 20, BODY, HIVE)

takeaway(s, "Read the axis numbers on the snapshots: the last panel is only 2 × 10⁻¹⁰ wide.")

# =====================================================================
# 17  All 40 runs met the target
# =====================================================================
s = new(17)

stat(s, "40 / 40", "runs finished below 10⁻⁶", 1.12, 2.92, 5.80, 1.72, AMBER)
stat(s, "5.476 × 10⁻¹³", "median final value", 7.10, 2.92, 5.80, 1.72, WHITE)
stat(s, "1.034 × 10⁻¹¹", "worst final value", 13.08, 2.92, 5.80, 1.72, WHITE)

card(s, 1.12, 4.92, 17.76, 4.62, WHITE)
pic(s, "11_reliability_40_runs.png", 1.26, 5.00, 17.48, 4.46)

takeaway(s, "Evidence of reliability on this problem — not a guarantee for every possible run.")

# =====================================================================
# 18  Bees next to genes
# =====================================================================
s = new(18)
deco(s, "hive_flowers", 17.00, -1.15, 3.6, 3.3)

table(s, ["", "Genetic Algorithm", "Artificial Bee Colony"], COMPARISON,
      1.12, 2.88, [3.95, 2.90, 3.30], 0.84, 22, highlight=2)

card(s, 1.12, 8.00, 9.78, 1.48, AMBER)
text(s, "Same reliability, three times closer,\none parameter to tune instead of five.",
     1.45, 8.26, 9.2, 1.0, 24, BODYB, BROWN, bold=True)

card(s, 11.25, 2.88, 7.63, 6.60, WHITE)
text(s, "The same evaluation budget, spent differently",
     11.50, 3.06, 7.1, 0.4, 20, BODYB, HIVE, bold=True)
pic(s, "13_budget_comparison.png", 11.40, 3.60, 7.33, 3.40)
text(s, "The Genetic Algorithm needed a crossover rate, a mutation rate, a mutation-step\n"
        "rule, elitism and a duplicate cap. ABC needed one number: the limit.",
     11.50, 7.30, 7.1, 1.9, 22, BODY, BROWN)

takeaway(s, "One function, two dimensions, 40 seeds. Evidence — not a claim about every problem.")

# =====================================================================
# 19  What we learned
# =====================================================================
s = new(19)
deco(s, "honeycomb", -3.05, 8.45, 6.60, 4.60)
deco(s, "hanging_hive", 16.80, 2.30, 2.85, 4.85)
deco(s, "bee", 17.40, 0.40, 1.10, 0.98)

for k, (title, bodytext) in enumerate([
        ("Three simple roles are enough",
         "Repeated 100 times, they reach 13 decimal places."),
        ("The step size comes free",
         "It is the distance between two bees, so it shrinks."),
        ("Test it, do not trust the diagram",
         "The scout phase never ran once on this problem."),
        ("Say what you did not test",
         "One function, two dimensions, 40 seeds.")]):
    x = 1.12 + (k % 2) * 7.75
    y = 2.95 + (k // 2) * 2.05
    bullet(s, f"0{k + 1}", title, bodytext, x, y, 7.70,
           [AMBER, DEEP, HIVE, BROWN][k], 27, 21)

text(s, "Thank you.  Questions?", 1.12, 7.55, 14.0, 1.40, 68, DISPLAY, BROWN)
text(s, "   ·   ".join(IDS), 1.12, 9.05, 14.0, 0.5, 22, BODY, HIVE)

# =====================================================================
# Save
# =====================================================================
prs.core_properties.title = ("Finding the Global Minimum of the 2-D Ackley Function "
                             "using the Artificial Bee Colony Algorithm")
prs.core_properties.subject = f"{COURSE} | {DATE}"
prs.core_properties.author = ", ".join(IDS)
prs.core_properties.keywords = ("Ackley, artificial bee colony, ABC, swarm intelligence, "
                                "roulette selection, KUET")
prs.core_properties.comments = ("Every number was reproduced from the notebook by "
                                "review/experiments.py. Illustrations are the original "
                                "template artwork.")

out = ROOT / "Ackley_ABC_Presentation.pptx"
prs.save(out)


# =====================================================================
# The speaking script
# =====================================================================
total = sum(d["seconds"] for d in SLIDES)
per_speaker = max(1, round(len(SLIDES) / len(IDS)))
chunks, start = [], 1
for n, student in enumerate(IDS):
    end = min(len(SLIDES), start + per_speaker - 1) if n < len(IDS) - 1 else len(SLIDES)
    if start <= end:
        chunks.append(f"{student}: slides {start}–{end}")
    start = end + 1

md = [
 "# Easy Presentation Script — Ackley Function with the Artificial Bee Colony",
 f"\n**{len(SLIDES)} slides · planned speaking time {total // 60} min {total % 60} sec "
 f"· presentation date {DATE}**",
 "\nThe short text belongs on the slides. The SAY paragraphs are also in the PowerPoint "
 "speaker notes. Speak naturally and point at the diagrams. Allow about 19–21 minutes "
 "including pauses and speaker changes.",
 f"\n**Six-speaker option:** {'; '.join(chunks)}. Adjust within the group if needed.",
 "\n**Pronunciation:** ABC = “A B C”; Ackley = “ACK-lee”; Karaboga = "
 "“kah-rah-BOW-ah”; phi = “fy”; 10⁻¹³ = “ten to the "
 "minus thirteen”; onlooker = “ON-look-er”; exploitation = “using what you "
 "already know”; exploration = “looking somewhere new”.",
]
for i, d in enumerate(SLIDES, 1):
    md += [f"\n---\n\n## Slide {i} — {d['title']}\n", "**ON THE SLIDE**\n"]
    md += ["- " + line for line in d["screen"]]
    md += [f"\n**FIGURE / DELIVERY:** {d['figure']}",
           f"\n**SAY**\n\n> {d['say']}",
           f"\n**TIME:** {d['seconds']} seconds."]

md += ["\n---\n\n## Short answers for teacher questions\n"]
for question, answer in QA:
    md += [f"### {question}\n\n{answer}\n"]

md += [
 "\n---\n\n## Where every number comes from\n",
 "All results were produced by `Ackley_ABC_Solution.ipynb` and independently re-run by "
 "`review/experiments.py`, which writes `review/results.json`. The main run uses seed 7.\n",
 "- Main run: x₁ = +2.278e-13, x₂ = −3.938e-14, f = 6.541e-13, best at cycle 99 "
 "of 100, distance 2.312e-13, 5,025 evaluations, 0 scouts.",
 "- 40 seeds (0–39): best 2.887e-14, median 5.476e-13, worst 1.034e-11, 40/40 below 1e-6, "
 "40/40 in the correct valley.",
 "- Random search, same 5,025-evaluation budget, same seed: f = 0.7693.",
 "- Genetic Algorithm, same 40 seeds, 5,050 evaluations: best 1.106e-13, median 1.636e-12, "
 "worst 3.137e-11, 40/40 below 1e-6.",
 "- Cycle-0 wheel: richest slice 9.56 % (source #7, f = 3.19), poorest 3.03 % (source #3, "
 "f = 12.23), ratio 3.16.",
 "- The wheel becomes exactly uniform (every slice 4.00 %) from cycle 62 onwards.",
 "- The largest trial counter in the main run reached 23, against a limit of 50, so no scout "
 "was ever sent.",
 "\n## Things to be careful about when answering\n",
 "- 5,025 is the number of Ackley calls that actually happened. 5,125 is the worst-case budget, "
 "which assumes a scout flies every cycle; none did. The slides quote the actual count.",
 "- “Success” means f < 1e-6. “Correct valley” means the distance to (0,0) is "
 "below 0.5. Neither means the answer is exactly zero.",
 "- Floating-point numbers do not have unlimited precision. Quote the measured errors rather "
 "than claiming an unlimited number of decimals.",
 "- The claim that the scout phase does not matter applies to **this** problem only. We did not "
 "test a harder or higher-dimensional function.",
 "- The comparison with the Genetic Algorithm is one problem at one budget. It is evidence, not "
 "proof that ABC is generally better.",
 "- The onlooker wheel going flat is real and measured, but it costs no measurable accuracy "
 "here; do not present it as a bug we fixed.",
 "\n## Rehearsal and slide use\n",
 "Use Presenter View for the notes. Slides 6 and 8 reveal their blocks on click; the PDF and the "
 "screenshots show the completed state. Rehearse slides 10, 13 and 14 most carefully, because "
 "they carry the reasoning. For a strict 18-minute slot, shorten the formula explanation on "
 "slide 3 and the ablation discussion on slide 14.",
]
(ROOT / "Ackley_ABC_Presentation_Script.md").write_text("\n".join(md) + "\n", encoding="utf-8")

print(json.dumps({
    "deck": str(out),
    "slides": len(prs.slides),
    "planned_seconds": total,
    "notes_words": sum(len(d["say"].split()) for d in SLIDES),
    "components_reused": sorted(COMPONENTS),
}, indent=2))
