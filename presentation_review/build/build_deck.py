from pathlib import Path
from copy import deepcopy
import sys, json
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from content import SLIDES, QA, DATE

ROOT=Path(__file__).resolve().parents[2]
REVIEW=ROOT/'presentation_review'
ASSETS=REVIEW/'assets'
NAME='Blue and Red Illustrated Genetic Biology Group Project Presentation.pptx'
prs=Presentation(ROOT/'presentation_backup'/NAME)
DNA=deepcopy(prs.slides[0].shapes[0]._element)
DNA_RELS={rid:(prs.slides[0].part.rels[rid].target_part,prs.slides[0].part.rels[rid].reltype) for rid in ['rId2','rId3']}
for sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)
BG='F3FAF9'; NAVY='163D50'; RED='B82E37'; TEAL='237B83'; INK='233C49'; MUTED='5B747D'; LINE='CFDEDF'; WHITE='FFFFFF'; PALE='E5F1F0'; PINK='FAE9E7'
FONT='Calibri'; DISPLAY='Arial'
I=Inches

def color(v): return RGBColor.from_string(v)
def box(s,x,y,w,h,fill=WHITE,line=None,radius=False,name=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,I(x),I(y),I(w),I(h))
    sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
    if line:sh.line.color.rgb=color(line)
    else:sh.line.fill.background()
    if radius:sh.adjustments[0]=0.08
    if name:sh.name=name
    return sh

def text(s,txt,x,y,w,h,size=30,bold=False,c=INK,align=None,font=FONT,name=None):
    sh=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=sh.text_frame
    tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c)
        p.space_after=Pt(6);p.space_before=Pt(0)
        if align is not None:p.alignment=align
    if name:sh.name=name
    return sh

def line(s,x1,y1,x2,y2,c=LINE,width=1.3):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2));sh.line.color.rgb=color(c);sh.line.width=Pt(width);return sh

def arrow(s,x,y,w=.5,h=.35,c=TEAL):
    sh=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,I(x),I(y),I(w),I(h));sh.fill.solid();sh.fill.fore_color.rgb=color(c);sh.line.fill.background();return sh

def motif(s,x,y,w,h):
    el=deepcopy(DNA)
    for node in el.iter():
        rid=node.get(qn('r:embed'))
        if rid in DNA_RELS:
            part,typ=DNA_RELS[rid];node.set(qn('r:embed'),s.part.relate_to(part,typ))
    s.shapes._spTree.insert_element_before(el,'p:extLst');sh=s.shapes[-1]
    sh.left=I(x);sh.top=I(y);sh.width=I(w);sh.height=I(h)
    sh._element.xpath('.//p:cNvPr')[0].set('id',str(s.shapes._next_shape_id))
    sh.name='Original template DNA illustration'
    return sh

def pic(s,name,x,y,w,h):
    p=ROOT/'presentation_figures'/name
    iw,ih=Image.open(p).size;r=min(w/iw,h/ih);pw=iw*r;ph=ih*r
    sh=s.shapes.add_picture(str(p),I(x+(w-pw)/2),I(y+(h-ph)/2),width=I(pw),height=I(ph));sh.name=name;return sh

def equation(s,formula,x,y,w,h,key):
    fig=plt.figure(figsize=(16,2));fig.patch.set_alpha(0)
    fig.text(.5,.5,'$'+formula+'$',ha='center',va='center',fontsize=27,color='#'+NAVY)
    p=ASSETS/(key+'.png');fig.savefig(p,dpi=230,bbox_inches='tight',pad_inches=.08,transparent=True);plt.close(fig)
    iw,ih=Image.open(p).size;r=min(w/iw,h/ih)
    sh=s.shapes.add_picture(str(p),I(x+(w-iw*r)/2),I(y+(h-ih*r)/2),width=I(iw*r),height=I(ih*r));sh.name=key
    sh._element.xpath('.//p:cNvPr')[0].set('descr',formula)
    return sh

def label(s,txt,x,y,w=8,c=TEAL):return text(s,txt,x,y,w,.35,18,True,c)
def takeaway(s,txt,y=9.62,c=NAVY):
    box(s,.8,y,18.4,.65,fill=c,radius=True)
    text(s,txt,1.05,y+.13,17.9,.4,24,True,WHITE)
def item(s,title,body,x,y,w=8.3,accent=TEAL):
    box(s,x,y+.02,.06,1.18,accent)
    text(s,title,x+.28,y,w-.4,.5,29,True,NAVY)
    text(s,body,x+.28,y+.57,w-.4,.85,25,c=MUTED)
def stat(s,value,desc,x,y,w=5.8,fill=WHITE):
    box(s,x,y,w,1.55,fill,radius=True)
    text(s,value,x+.25,y+.15,w-.5,.7,42,True,RED)
    text(s,desc,x+.25,y+.98,w-.5,.42,23,c=MUTED)
def table(s,headers,rows,x,y,widths,rh=.8,fs=26,highlight=None):
    sh=s.shapes.add_table(len(rows)+1,len(headers),I(x),I(y),I(sum(widths)),I(rh*(len(rows)+1)))
    t=sh.table
    for j,w in enumerate(widths):t.columns[j].width=I(w)
    for i,row in enumerate([headers]+rows):
        t.rows[i].height=I(rh)
        for j,val in enumerate(row):
            cell=t.cell(i,j);cell.text=str(val);cell.margin_left=I(.18);cell.margin_right=I(.14);cell.margin_top=I(.10);cell.margin_bottom=I(.05)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=color(NAVY if i==0 else (PALE if i==highlight else (WHITE if i%2 else 'EDF5F4')))
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(fs);p.font.bold=i==0 or i==highlight;p.font.color.rgb=color(WHITE if i==0 else INK)
            if i>0 and str(val)=='2.58':
                for p in cell.text_frame.paragraphs:p.font.bold=True;p.font.color.rgb=color(RED)
    return sh

def new(n):
    data=SLIDES[n-1];s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(BG)
    for sh in list(s.shapes):s.shapes._spTree.remove(sh._element)
    box(s,0,0,20,.10,RED)
    label(s,data['section'],.8,.4,17)
    if n!=1:text(s,data['title'],.8,1.0,18.4,.95,45,True,NAVY,font=DISPLAY,name='Slide title')
    line(s,.8,10.55,19.2,10.55)
    text(s,DATE,.8,10.72,5,.3,17,c=MUTED,name='Presentation date')
    text(s,'ACKLEY × GENETIC ALGORITHM',6,10.72,8,.3,15,c=MUTED,align=PP_ALIGN.CENTER)
    text(s,f'{n:02d} / 18',16.6,10.69,2.6,.4,20,True,NAVY,PP_ALIGN.RIGHT,name='Slide number')
    s.notes_slide.notes_text_frame.text=f"SLIDE {n} — {data['title']}\nTarget: {data['seconds']} seconds\n\n{data['say']}\n\nVISUAL / DELIVERY\n{data['figure']}\n\nSource: Ackley_GA_Solution.ipynb; numerical results independently reproduced."
    return s

s=new(1)
motif(s,16.8,1.5,2.0,7.7)
label(s,'GENETIC ALGORITHM · GROUP PRESENTATION',1.0,1.55,14,RED)
text(s,'Finding the\nglobal minimum',1.0,2.25,15,2.35,66,True,NAVY,font=DISPLAY)
text(s,'The 2-D Ackley function using a Genetic Algorithm',1.0,4.95,15,.7,31,c=TEAL)
line(s,1.0,6.1,15.2,6.1,LINE,2)
label(s,'SUBMITTED TO',1,6.55,13)
text(s,'Dr. Al-Mahmud',1,7.03,13,.57,32,True)
text(s,'Professor, Department of CSE, KUET',1,7.63,14,.45,26,c=MUTED)
label(s,'SUBMITTED BY · STUDENT IDs',1,8.55,13)
text(s,'2107001  ·  2107004  ·  2107009\n2107015  ·  2107024  ·  2107047',1,9.02,14,1.0,27,c=INK)

s=new(2)
item(s,'The task','Choose x₁ and x₂ to make f as small as possible.',.9,2.65,9)
item(s,'The search range','Both coordinates stay between −5 and +5.',.9,4.35,9)
item(s,'The known answer','x* = (0, 0)   and   f(x*) = 0',.9,6.05,9,RED)
box(s,11,2.4,7.4,6.5,WHITE,radius=True)
for k in range(6):
    xx=12.1+k*.95; yy=3.2+k*.9
    line(s,xx,3.2,xx,7.7,'E3ECEC',.7);line(s,12.1,yy,16.85,yy,'E3ECEC',.7)
line(s,12.1,5.45,17.3,5.45,MUTED,2);line(s,14.48,3.0,14.48,8.0,MUTED,2)
dot=s.shapes.add_shape(MSO_SHAPE.OVAL,I(14.36),I(5.33),I(.24),I(.24));dot.fill.solid();dot.fill.fore_color.rgb=color(RED);dot.line.fill.background()
text(s,'(0, 0)',14.8,5.66,2,.5,27,True,RED)
text(s,'−5',11.9,7.9,1,.4,21,c=MUTED);text(s,'+5',16.5,7.9,1,.4,21,c=MUTED)
text(s,'x₁',17.4,5.24,.6,.5,27,True);text(s,'x₂',14.2,2.58,.8,.5,27,True)
takeaway(s,'The search evaluates candidate points; it is not given the answer coordinates.')

s=new(3)
label(s,'START WITH n VARIABLES',1,2.2)
box(s,.9,2.72,18.2,1.65,WHITE,radius=True)
equation(s,r'f(\mathbf{x})=-a\exp\!\left(-b\sqrt{\frac{1}{n}\sum_{i=1}^{n}x_i^2}\right)-\exp\!\left(\frac{1}{n}\sum_{i=1}^{n}\cos(cx_i)\right)+a+e',1.2,2.9,17.6,1.22,'general_ackley')
for txt,x in [('Set n = 2',1.5),('Expand the two sums',7.15),('Use a=20, b=0.2, c=2π',12.8)]:text(s,txt,x,4.82,5.1,.7,28,True,TEAL,PP_ALIGN.CENTER)
arrow(s,6.65,4.98);arrow(s,12.25,4.98)
label(s,'THE FUNCTION IN OUR CODE',1,5.75)
box(s,.9,6.25,18.2,2.1,WHITE,radius=True)
equation(s,r'f(x_1,x_2)=-20\exp\!\left(-0.2\sqrt{\frac{x_1^2+x_2^2}{2}}\right)-\exp\!\left(\frac{\cos(2\pi x_1)+\cos(2\pi x_2)}{2}\right)+20+e',1.2,6.65,17.6,1.2,'two_dimensional_ackley')
takeaway(s,'At the origin:   −20 − e + 20 + e = 0')

s=new(4)
box(s,.8,2.25,12.2,5.5,WHITE,radius=True);pic(s,'01_ackley_landscape.png',.95,2.4,11.9,5.2)
pic(s,'02_ackley_slice.png',1.0,7.8,11.8,2.45)
item(s,'A wide funnel','The smooth part points toward the centre.',13.65,2.65,5.4)
item(s,'Many local dips','The cosine part adds small traps.',13.65,4.55,5.4,RED)
item(s,'Expensive to scan','A 10⁻⁶ grid needs about 10¹⁴ points.',13.65,6.45,5.4)
text(s,'A downhill method can stop\nin a local minimum.',13.93,8.6,5.0,1.0,26,True,RED)

s=new(5)
label(s,'BIOLOGY → OUR PROGRAM',1,2.4)
table(s,['Term','Meaning'],[['Population','50 candidate points'],['Chromosome','One pair [x₁, x₂]'],['Gene','One coordinate'],['Fitness','Selection weight']],1,3.03,[3.25,5.6],.9,28)
text(s,'Evaluate → select → mix\n→ mutate → repeat',1.2,8.0,8.7,1.25,31,True,TEAL)
box(s,10.4,2.27,8.65,7.0,WHITE,radius=True);pic(s,'03_initial_population.png',10.65,2.4,8.1,6.7)
takeaway(s,'A population explores many points. Better candidates get more chances to reproduce.')

s=new(6)
label(s,'VALUE ENCODING',1,2.4)
for x,v,g,c in [(3.3,'x₁','GENE 1',TEAL),(10.5,'x₂','GENE 2',RED)]:
    label(s,g,x,3.12,6,c);box(s,x,3.65,6.2,2.3,WHITE,radius=True);text(s,v,x,3.92,6.2,1.6,85,True,c,PP_ALIGN.CENTER)
text(s,'One chromosome',6.25,6.24,7.5,.6,32,True,NAVY,PP_ALIGN.CENTER)
table(s,['Example A','Example B'],[['[−4.0, 2.5]','[−1.5, −1.0]']],2.5,7.2,[7.5,7.5],.73,28)
takeaway(s,'Each gene stays in [−5, +5]. Real-valued encoding still has finite computer precision.')

s=new(7)
label(s,'ASSIGNMENT SETTINGS',1,2.4,8)
label(s,'OUR IMPLEMENTATION CHOICES',10.5,2.4,9,RED)
table(s,['Parameter','Value'],[['Population','50'],['Max. generations','100'],['Crossover','80%'],['Mutation','5% per gene'],['Elitism','1']],1,3.0,[5.1,3.3],.79,28)
text(s,'Value encoding · roulette selection\nOne-point crossover',1.15,8.2,8.6,1.1,26,True,TEAL)
table(s,['Choice','Value'],[['Mutation type','Gaussian'],['Step size σ','clip(f_best/4, 10⁻¹², 2)'],['Exact-copy limit','15 of 50'],['Tolerance / patience','10⁻¹⁰ / 20 rounds'],['Main random seed','7']],10.5,3,[4.3,4.2],.79,25)
text(s,'Probability is fixed;\nstep size adapts to the result.',10.65,8.2,8.3,1.1,27,True,RED)

s=new(8)
def flowblock(label_,body,x,y,w,h=1.35,fill=WHITE,name=None):
    sh=box(s,x,y,w,h,fill,radius=True,name=name)
    t=text(s,label_,x+.22,y+.18,w-.44,.5,29,True,NAVY,name=(name+'_title') if name else None)
    if body:text(s,body,x+.22,y+.77,w-.44,.53,22,c=MUTED,name=(name+'_body') if name else None)
    return sh
flowblock('1  Initialise','50 random points → evaluate → sort',1,2.45,8.35)
flowblock('2  Check stopping rules','If yes: report best and finish',10.5,2.45,8.5,fill=PALE)
arrow(s,9.55,2.95,.65,.42)
line(s,14.75,3.8,14.75,4.47,TEAL,2);text(s,'NO: CONTINUE',15,3.97,3,.35,17,True,TEAL)
flowblock('3  Set σ and probabilities','Then select 50 parents by roulette',10.5,4.65,8.5,name='Reveal_1')
flowblock('4  Make 50 children','Crossover → mutate → clip → evaluate',1,4.65,8.35,name='Reveal_2')
ar=arrow(s,9.55,5.1,.65,.42);ar.rotation=180
line(s,5.15,6.0,5.15,6.55,TEAL,2)
flowblock('5  Choose the next 50','1 elite + 49 seats, with the copy limit',1,6.8,8.35,fill=PALE,name='Reveal_3')
flowblock('6  Repeat','Use the survivors as the next population',10.5,6.8,8.5,name='Reveal_4')
arrow(s,9.55,7.25,.65,.42)
text(s,'↺  Return to the stopping check',10.7,8.62,8.1,.5,28,True,TEAL)
takeaway(s,'50 parents + 50 children → preserve the elite and fill the next population of 50.')

s=new(9)
text(s,'Lower objective → higher fitness',1,2.3,9.8,.55,31,True,TEAL)
box(s,.9,3.05,10.3,2.25,WHITE,radius=True)
equation(s,r'F_i=(f_{\max}-f_i)+\delta,\qquad p_i=\frac{F_i}{\sum_j F_j}',1.15,3.25,9.8,.85,'fitness')
text(s,'δ = 0.10(f_max − f_min) + 10⁻¹²',1.25,4.5,9.5,.5,25,c=MUTED)
table(s,['Chromosome','f','Probability'],[['#7 (best)','3.19','4.83%'],['#46','3.79','4.57%'],['#14','4.87','4.10%'],['#45 (worst)','13.39','0.44%']],1,5.77,[4.05,2.1,4.05],.68,25)
box(s,11.65,2.2,7.55,7.4,WHITE,radius=True);pic(s,'04_roulette_pie_gen0.png',11.75,2.3,7.35,7.12)
takeaway(s,'The code uses 50 probabilities. The chart groups the remaining 40 into one grey slice.')

s=new(10)
label(s,'PARENTS',1.35,2.65,7);label(s,'CHILDREN',12,2.65,6,RED)
def genes(s,x,y,vals,colors,name=None):
    for j,(v,c) in enumerate(zip(vals,colors)):
        box(s,x+j*2.75,y,2.55,1.5,c,radius=True,name=f'{name}_{j}' if name else None)
        text(s,v,x+j*2.75,y+.35,2.55,.8,40,True,WHITE,PP_ALIGN.CENTER,name=f'{name}_text_{j}' if name else None)
genes(s,1.4,3.45,['−4.0','2.5'],[TEAL,TEAL]);genes(s,1.4,5.65,['−1.5','−1.0'],[RED,RED])
genes(s,12.0,3.45,['−4.0','−1.0'],[TEAL,RED],'Reveal_1');genes(s,12.0,5.65,['−1.5','2.5'],[RED,TEAL],'Reveal_2')
arrow(s,8.25,4.75,2.1,.8)
text(s,'80%\nswap tails',7.25,3.35,4.1,1.2,30,True,NAVY,PP_ALIGN.CENTER)
text(s,'x₁',2.15,7.45,1.5,.5,25,True,TEAL);text(s,'x₂',4.85,7.45,1.5,.5,25,True,RED)
text(s,'20%: children are unchanged copies of the parents.',1.4,8.5,17,.65,31,c=MUTED)
takeaway(s,'Only one legal cut exists. Crossover re-pairs values that are already in the population.')

s=new(11)
item(s,'1  Check each gene','Each coordinate has its own 5% chance.',1,2.5,8.8)
item(s,'2  Add a random change','xᵢ ← xᵢ + N(0, σ²)',1,4.4,8.8)
item(s,'3  Repair the boundary','Clip the changed coordinate to [−5, +5].',1,6.3,8.8)
box(s,10.6,2.4,8.4,6.55,WHITE,radius=True)
label(s,'ILLUSTRATION · NOT A RECORDED DRAW',11,2.85,7.7)
text(s,'[−1.5, −1.0]',11.1,3.75,7.2,.9,43,True,NAVY,PP_ALIGN.CENTER)
text(s,'↓  first gene changes',11.3,4.8,7,.7,28,c=MUTED,align=PP_ALIGN.CENTER)
text(s,'[−0.5, −1.0]',11.1,5.75,7.2,.9,43,True,RED,PP_ALIGN.CENTER)
text(s,'Boundary example: 5.3 → 5.0',11.2,7.55,7.3,.7,27,True,TEAL,PP_ALIGN.CENTER)
takeaway(s,'The 5% probability stays fixed. The size of the change is controlled by σ.')

s=new(12)
box(s,.9,2.3,7.7,2.2,PINK,radius=True)
text(s,'Timed schedule',1.2,2.62,7.1,.6,32,True,RED)
text(s,'Steps shrink even when\nthe run is still stuck.',1.2,3.22,7.1,1.13,29)
box(s,.9,4.92,7.7,2.35,WHITE,radius=True)
text(s,'Our adaptive rule',1.2,5.24,7.1,.5,29,True,TEAL)
text(s,'σ = clip(f_best / 4, 10⁻¹², 2)',1.2,6.0,7.1,.8,29,True,NAVY)
table(s,['Best value f','Step size σ'],[['5','1.25'],['2.6','0.65'],['10⁻³','2.5 × 10⁻⁴'],['10⁻¹²','10⁻¹² (floor)']],9.25,2.3,[4.8,4.7],.9,30)
text(s,'Near the origin:  f ≈ 4s',9.5,7.22,9.1,.7,32,True,TEAL)
equation(s,r's=\sqrt{(x_1^2+x_2^2)/2}',9.45,8.0,8.6,.85,'local_scale')
takeaway(s,'f/4 estimates a local coordinate scale for Ackley; it is not the exact distance to the answer.')

s=new(13)
stat(s,'50','parents',1,2.5,5.5);text(s,'+',6.73,2.83,1,.9,46,True,TEAL)
stat(s,'50','children',8,2.5,5.5);text(s,'=',13.75,2.83,1,.9,46,True,TEAL)
stat(s,'100','candidates',15,2.5,4)
box(s,1,4.75,6.0,3.62,NAVY,radius=True,name='Reveal_1')
text(s,'1',1.35,5.03,5.3,1.12,68,True,WHITE,name='Reveal_1_text')
text(s,'Keep the elite',1.35,6.48,5.3,.6,31,True,WHITE,name='Reveal_1_title')
text(s,'Best parent is protected.',1.35,7.28,5.3,.5,24,c='CCE5E5',name='Reveal_1_body')
box(s,7.6,4.75,11.4,3.62,WHITE,radius=True,name='Reveal_2')
text(s,'49 more seats',8,5.17,10.6,.75,39,True,TEAL,name='Reveal_2_title')
text(s,'Best remaining candidates, with ≤15 exact copies.\nIf needed, fill empty seats with fresh random points.',8,6.36,10.55,1.35,29,name='Reveal_2_body')
takeaway(s,'Elitism protects the best-so-far value. The copy limit keeps more distinct candidates.')

s=new(14)
box(s,.9,2.3,8.65,1.65,WHITE,radius=True)
text(s,'100 generations',1.2,2.53,8.0,.7,39,True,RED);text(s,'maximum number of rounds',1.2,3.3,8,.43,23,c=MUTED)
text(s,'OR',9.65,2.83,1.1,.6,27,True,TEAL,PP_ALIGN.CENTER)
box(s,10.8,2.3,8.3,1.65,WHITE,radius=True)
text(s,'20 small improvements in a row',11.08,2.56,7.8,.6,28,True,NAVY);text(s,'each improvement < 10⁻¹⁰',11.08,3.31,7.8,.43,23,c=MUTED)
box(s,.9,4.38,18.2,4.93,WHITE,radius=True);pic(s,'05_convergence_and_diversity.png',1.08,4.53,17.85,4.54)
text(s,'Note: the plotted σ label is shorthand; the code clips σ to [10⁻¹², 2].',1.18,9.17,17.6,.35,18,c=MUTED)
takeaway(s,'Seed 7: f ≈ 0.00587 at generation 20; the best result arrives at generation 100.')

s=new(15)
stat(s,'1.088 × 10⁻¹²','final objective value · seed 7',.9,2.2,6.0)
stat(s,'100','generation of the best result',7,2.2,5.9)
stat(s,'5,050','function evaluations',13,2.2,6.1)
text(s,'x₁ = −3.436 × 10⁻¹³     x₂ = −1.730 × 10⁻¹³     distance = 3.847 × 10⁻¹³',1.0,4.08,18,.6,29,True,TEAL)
box(s,.9,4.92,18.2,4.1,WHITE,radius=True);pic(s,'06_best_solution_path.png',1.05,4.99,17.9,3.92)
text(s,'Both coordinates have absolute error < 10⁻¹². The result is approximate.',1.15,9.18,17.7,.4,23,c=MUTED)
takeaway(s,'Random search with the same seed and 5,050 evaluations: f = 0.7693. This is one comparison.')

s=new(16)
text(s,'Same 40 seeds in every row',1,2.2,18,.55,29,c=MUTED)
table(s,['Step rule','Copy cap','Median f','Worst f','In valley','Distinct'],[
['Timed','None','5.52 × 10⁻⁹','2.58','39 / 40','3'],
['Timed','15','6.76 × 10⁻⁹','2.58','39 / 40','9'],
['Adaptive','None','5.00 × 10⁻¹³','2.31 × 10⁻¹¹','40 / 40','8'],
['Adaptive','15','1.64 × 10⁻¹²','3.14 × 10⁻¹¹','40 / 40','13']],.9,3.05,[3.1,2.1,3.65,3.7,2.95,2.7],.91,26,4)
text(s,'In valley: distance to (0, 0) < 0.5. Distinct: median number of unique chromosomes\nover the second half of each run, then the median across runs.',1.1,7.78,17.9,.9,21,c=MUTED)
box(s,.95,9.1,8.8,1.15,PALE,radius=True);text(s,'Adaptive step → smaller errors',1.2,9.42,8.3,.52,28,True,TEAL)
box(s,10.15,9.1,8.85,1.15,PINK,radius=True);text(s,'Copy limit → more distinct points',10.4,9.42,8.3,.52,28,True,RED)

s=new(17)
stat(s,'40 / 40','passed f < 10⁻⁶ · seeds 0–39',.9,2.2,6)
stat(s,'1.636 × 10⁻¹²','median final value',7,2.2,6)
stat(s,'3.137 × 10⁻¹¹','worst final value',13.1,2.2,6)
box(s,.9,4.23,18.2,5.02,WHITE,radius=True);pic(s,'09_reliability_40_runs.png',1.05,4.4,17.9,4.71)
text(s,'Best f = 1.106 × 10⁻¹³. Points visually overlap the optimum; they are approximate.',1.15,9.22,17.8,.38,21,c=MUTED)
takeaway(s,'All tested seeds met the target. This is evidence of reliability, not a guarantee for every run.')

s=new(18)
motif(s,17.25,2.75,1.4,5.85)
text(s,'A good result needs\na well-tested search.',1,2.35,15.5,1.65,46,True,RED,font=DISPLAY)
for title_,body_,x,y in [('Select + mix','Better parents; useful coordinate combinations.',1,4.75),('Mutate + adapt','New values; steps that shrink with progress.',8.8,4.75),('Preserve + diversify','Keep the best; limit exact copies.',1,6.65),('Check many seeds','40/40 below the target in our tests.',8.8,6.65)]:item(s,title_,body_,x,y,7.45)
text(s,'Thank you. Questions?',1,9.05,15.5,.9,42,True,NAVY)

# Static dates are intentional: this is the presentation date, not the viewing date.
prs.core_properties.title='Finding the Global Minimum of the 2-D Ackley Function using a Genetic Algorithm'
prs.core_properties.subject='CSE 4111 — Machine Learning | 14 September 2026'
prs.core_properties.author='2107001, 2107004, 2107009, 2107015, 2107024, 2107047'
prs.core_properties.keywords='Ackley, genetic algorithm, roulette selection, KUET'
prs.core_properties.comments='Content checked against the notebook and independently reproduced. Original template illustrations retained.'
out=ROOT/'Ackley_GA_Presentation_Updated.pptx';prs.save(out)

total=sum(d['seconds'] for d in SLIDES)
md=['# Easy Presentation Script — Ackley Function with a Genetic Algorithm',f'\n**18 slides · planned speaking time {total//60} min {total%60} sec · presentation date {DATE}**',
'\nThe short text belongs on the slides. The SAY paragraphs are also in the PowerPoint speaker notes. Speak naturally and point to the diagrams. Allow about 17–19 minutes including pauses and speaker changes.',
'\n**Six-speaker option:** IDs 2107001: slides 1–3; 2107004: 4–6; 2107009: 7–9; 2107015: 10–12; 2107024: 13–15; 2107047: 16–18. Adjust within the group if needed.',
'\n**Pronunciation:** GA = “G A”; Ackley = “ACK-lee”; sigma = “SIG-ma”; 10⁻¹² = “ten to the minus twelve”; elitism = “keep the best”; convergence = “the result settles down”.']
for i,d in enumerate(SLIDES,1):
    md += [f'\n---\n\n## Slide {i} — {d["title"]}\n', '**ON THE SLIDE**\n']+['- '+x for x in d['screen']]+[f'\n**FIGURE / DELIVERY:** {d["figure"]}',f'\n**SAY**\n\n> {d["say"]}',f'\n**TIME:** {d["seconds"]} seconds.']
md += ['\n---\n\n## Short answers for teacher questions\n']
for q,a in QA:md += [f'### {q}\n\n{a}\n']
md += ['\n## Notebook check and corrections\n',
'The numerical results were independently reproduced from the notebook code, including the main run and the four configurations with 40 seeds each. Verification data: `presentation_review/notebook_verification.json`. The notebook itself has not been edited.',
'\n- The exact adaptive rule includes clipping: `sigma = clip(best_value / 4, 1e-12, 2.0)`.',
'- The local scale `sqrt((x1²+x2²)/2)` is not the Euclidean distance. The step rule is Ackley-specific and uses knowledge of its local shape.',
'- Floating-point numbers do not have unlimited precision. Use the measured absolute errors instead of claiming an unlimited number of decimals.',
'- The final run found its best at generation 100. Generation 20 only reached about 0.00587.',
'- Success is defined as `f < 1e-6`. Correct valley is a separate check: distance to the origin below 0.5. Neither means exact equality with zero.',
'- The timed versions reached the correct valley in 39 of 40 runs: exactly 97.5%, shown as 98% by the notebook formatting. The slides use counts to avoid rounding confusion.',
'- The copy limit counts exact chromosomes. Different chromosomes can still be very close. Fresh random fill points are allowed and are also a source of new coordinate values.',
'- Elitism guarantees that the best-so-far value cannot worsen; it does not guarantee that the population average cannot worsen.',
'- The seed-7 run used exactly 5,050 objective-point evaluations, verified by counting calls. Random fill points would add evaluations in a run that needs them.',
'- The baseline result is one same-budget comparison. Removed the broad “700 billion times more accurate” claim.',
'- Removed unsupported claims that cap values 10–25 had been systematically validated, and that both fixes are necessary for every successful run.',
'- The roulette pie groups the other 40 chromosomes for display; actual selection uses all 50 probabilities.',
'\n## Rehearsal and slide use\n',
'Use Presenter View for the notes. Slides 8, 10 and 13 have click reveals. The PDF and screenshots show their completed state. Rehearse slides 12 and 16 most carefully. Keep the pace near the suggested times, and pause after explaining each figure. For a strict 15-minute slot, shorten the formula explanation and the comparison table discussion.']
(ROOT/'Ackley_GA_Presentation_Script.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
print(json.dumps({'deck':str(out),'slides':len(prs.slides),'planned_seconds':total,'notes_words':sum(len(d['say'].split()) for d in SLIDES)},indent=2))
